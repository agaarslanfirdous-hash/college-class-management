#!/usr/bin/env python3
"""Rebuild the deck around AI + IoT governance, with PDN as one layer."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFilter, ImageEnhance, ImageOps, ImageStat
import os


BASE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(BASE, "assets")
OUT = os.path.join(BASE, "Arsalan_Firdous_AI_IoT_Paddy_Irrigation_UNNATI_v2.pptx")
USER_ASSETS = "/Users/arsalanfirdous/.cursor/projects/Users-arsalanfirdous-Downloads-Aaaag/assets"

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(13, 43, 58)
TEAL = RGBColor(0, 107, 122)
TEAL_DARK = RGBColor(10, 61, 74)
CYAN = RGBColor(39, 177, 190)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(57, 67, 75)
MUTED = RGBColor(88, 101, 109)
GREEN = RGBColor(28, 124, 86)
ORANGE = RGBColor(196, 107, 46)
LIGHT = RGBColor(238, 250, 251)

TOPIC = "AI + IoT Integrated Paddy Irrigation\nDecision & Governance Platform"
SUBTOPIC = "Multimodal sensing, predictive irrigation, remote crop intelligence, and human-supervised water governance"


def set_run(run, size=18, bold=False, color=NAVY, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)


def make_card(src_path, out_path, max_h=80, pad=14):
    im = Image.open(src_path).convert("RGBA")
    ratio = max_h / im.height
    im = im.resize((int(im.width * ratio), max_h), Image.Resampling.LANCZOS)
    card = Image.new("RGBA", (im.width + pad * 2, im.height + pad * 2), (255, 255, 255, 245))
    card.paste(im, (pad, pad), im)
    card.save(out_path)


def build_backgrounds():
    ensure_dir(ASSETS)
    make_card(os.path.join(ASSETS, "skuast_logo.png"), os.path.join(ASSETS, "skuast_header.png"), 70, 14)
    make_card(os.path.join(ASSETS, "unnati_official.png"), os.path.join(ASSETS, "unnati_header.png"), 70, 14)

    bg_wave = os.path.join(ASSETS, "bg_wave.jpg")
    if os.path.exists(bg_wave):
        im = Image.open(bg_wave).convert("RGB").resize((1920, 1080), Image.Resampling.LANCZOS)
        im.save(os.path.join(ASSETS, "bg_title_soft.jpg"), quality=94)

    sources = {
        "bg_problem.jpg": "Screenshot_2026-08-25_at_6.58.45_PM-584ae831-2580-4d4f-a618-c6f2f1504e97.jpg",
        "bg_architecture.jpg": "Screenshot_2026-08-25_at_6.58.07_PM-fd0dec7c-6bfe-4d27-9b71-470e5370c446.jpg",
        "bg_ai.jpg": "Screenshot_2026-08-25_at_6.58.01_PM-849e6f80-98db-4d18-a98f-4e4211f70214.jpg",
        "bg_sensors.jpg": "Screenshot_2026-08-25_at_6.58.38_PM-ee7c5aae-e332-4aab-ac34-9d264eeb4fa6.jpg",
        "bg_infra.jpg": "Screenshot_2026-08-25_at_6.57.49_PM-2bb20946-2f65-456d-b356-3699c970778e.jpg",
        "bg_method.jpg": "Screenshot_2026-08-25_at_6.58.29_PM-8f716930-96c4-45e5-867b-4121dc16da4b.jpg",
    }

    for out_name, src_name in sources.items():
        src_path = os.path.join(USER_ASSETS, src_name)
        if not os.path.exists(src_path):
            continue
        im = Image.open(src_path).convert("RGB").resize((1920, 1080), Image.Resampling.LANCZOS)
        im = ImageEnhance.Color(im).enhance(0.92)
        im = ImageEnhance.Contrast(im).enhance(1.05)
        im = im.filter(ImageFilter.GaussianBlur(7))

        overlay = Image.new("RGBA", im.size, (6, 34, 44, 132))
        comp = Image.alpha_composite(im.convert("RGBA"), overlay)

        # Add a soft light band in the center to keep text readable.
        fade = Image.new("RGBA", im.size, (255, 255, 255, 0))
        for y in range(im.size[1]):
            alpha = 0
            if 120 <= y <= 940:
                dist = abs(y - 530) / 410.0
                alpha = max(0, int(130 * (1 - dist)))
            for x in range(im.size[0]):
                fade.putpixel((x, y), (255, 255, 255, alpha))
        comp = Image.alpha_composite(comp, fade)
        comp.convert("RGB").save(os.path.join(ASSETS, out_name), quality=92)


def add_bg(slide, prs, name):
    slide.shapes.add_picture(os.path.join(ASSETS, name), 0, 0, width=prs.slide_width, height=prs.slide_height)


def brand_header(slide):
    sku = os.path.join(ASSETS, "skuast_header.png")
    unnati = os.path.join(ASSETS, "unnati_header.png")
    if os.path.exists(sku):
        slide.shapes.add_picture(sku, Inches(0.32), Inches(0.14), height=Inches(0.58))
    if os.path.exists(unnati):
        slide.shapes.add_picture(unnati, Inches(9.45), Inches(0.14), height=Inches(0.58))


def add_footer(slide, prs, page, total):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.04), prs.slide_width, Inches(0.46))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL_DARK
    shape.line.fill.background()
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"Arsalan Firdous  ·  SKUAST-Kashmir  ·  UNNATI AI 2.0 Roadshow  ·  {page}/{total}"
    set_run(run, 11, False, WHITE)


def add_title(slide, text, top=0.9, size=28, color=TEAL_DARK):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.1), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size, True, color)


def add_panel(slide, x, y, w, h, fill=WHITE, alpha=235, line=RGBColor(184, 220, 224)):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    try:
        s.fill.transparency = max(0.0, min(1.0, 1 - alpha / 255))
    except Exception:
        pass
    s.line.color.rgb = line
    s.line.width = Pt(1)
    s.adjustments[0] = 0.08
    return s


def add_text(slide, x, y, w, h, lines, size=15, color=NAVY, bold=False, align=PP_ALIGN.LEFT, space_after=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return box


def bullet_panel(slide, x, y, w, h, title, bullets, title_color=TEAL, bullet_size=13):
    add_panel(slide, x, y, w, h)
    add_text(slide, x + 0.18, y + 0.12, w - 0.35, 0.35, [title], 15, title_color, True)
    add_text(slide, x + 0.18, y + 0.46, w - 0.35, h - 0.55, [f"•  {b}" for b in bullets], bullet_size, GRAY)


def add_stat(slide, x, y, w, h, value, label, ref):
    add_panel(slide, x, y, w, h, fill=WHITE, alpha=242)
    add_text(slide, x + 0.05, y + 0.15, w - 0.1, 0.48, [value], 26, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, x + 0.12, y + 0.78, w - 0.24, h - 0.85, [f"{label} {ref}"], 11, GRAY, False, PP_ALIGN.CENTER)


def add_arch_box(slide, x, y, w, h, title, lines, fill, title_color=WHITE):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.adjustments[0] = 0.08
    add_text(slide, x + 0.12, y + 0.12, w - 0.24, 0.32, [title], 14, title_color, True, PP_ALIGN.CENTER)
    add_text(slide, x + 0.14, y + 0.46, w - 0.28, h - 0.52, lines, 11, WHITE, False, PP_ALIGN.CENTER, 2)


def build_presentation():
    build_backgrounds()
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    total = 12

    refs = [
        "[1] India canal conveyance loss / hydrospatial modelling studies",
        "[2] PDN/CDN efficiency and CWC-aligned piped irrigation literature",
        "[3] Rice water use and deep percolation loss studies",
        "[4] Climate and irrigation constraints in rice systems",
        "[5] IoT-assisted AWD and automated rice irrigation studies",
        "[6] Multispectral / hyperspectral crop stress and disease monitoring literature",
        "[7] Precision irrigation and human-in-the-loop governance framing",
    ]

    # 1 title
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_title_soft.jpg")
    brand_header(s)
    add_panel(s, 1.45, 1.35, 10.45, 4.95, fill=WHITE, alpha=240)
    add_text(s, 1.8, 1.6, 9.75, 0.35, ["UNNATI AI 2.0 Roadshow"], 14, CYAN, True, PP_ALIGN.CENTER)
    add_text(s, 1.8, 2.05, 9.75, 1.1, [TOPIC], 26, TEAL_DARK, True, PP_ALIGN.CENTER)
    add_text(s, 1.8, 3.25, 9.75, 0.55, [SUBTOPIC], 13, TEAL, False, PP_ALIGN.CENTER)
    add_text(
        s, 1.8, 4.1, 9.75, 1.5,
        [
            "Presented by: Arsalan Firdous",
            "Sher-e-Kashmir University of Agricultural Sciences & Technology (SKUAST-Kashmir)",
            "Closed-loop agricultural water governance: Sense → Analyze → Decide → Actuate → Govern",
        ],
        15, GRAY, False, PP_ALIGN.CENTER
    )
    add_footer(s, prs, 1, total)

    # 2 problem snapshot
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_problem.jpg")
    brand_header(s)
    add_title(s, "Why Current Irrigation Approaches Fall Short")
    add_stat(s, 0.55, 1.75, 2.9, 2.05, "~45%", "water may be lost during conveyance", "[1]")
    add_stat(s, 3.65, 1.75, 2.9, 2.05, "60-83%", "applied paddy water may be lost to deep percolation", "[3]")
    add_stat(s, 6.75, 1.75, 2.9, 2.05, "20-50%", "water saving possible with better scheduling vs flooding", "[4][5]")
    add_stat(s, 9.85, 1.75, 2.9, 2.05, "13-20%", "extra saving reported when IoT improves AWD decisions", "[5]")
    bullet_panel(
        s, 0.55, 4.2, 12.2, 2.3, "The real gap is not only irrigation hardware",
        [
            "Water is delivered without continuous knowledge of crop stage, soil status, forecast, disease risk, or true field demand.",
            "Conventional basin flooding and fixed canal turns are not designed for dynamic, field-level decisions.",
            "Departments and farmers usually cannot verify allocation, crop stress, leakages, or waterlogging in real time.",
        ],
    )
    add_footer(s, prs, 2, total)

    # 3 why ai+iot
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_ai.jpg")
    brand_header(s)
    add_title(s, "Why the Core Must Be AI + IoT, Not Just Irrigation Replacement")
    bullet_panel(
        s, 0.6, 1.65, 4.0, 4.95, "AI + IoT value",
        [
            "Turns irrigation from schedule-based delivery into demand-based decision-making.",
            "Fuses live sensors, imagery, weather, and hydraulic data into one operational intelligence layer.",
            "Detects water stress, waterlogging, disease risk, pest risk, and infrastructure anomalies earlier.",
            "Supports both automation and human override instead of blind autonomy.",
        ],
    )
    bullet_panel(
        s, 4.72, 1.65, 4.0, 4.95, "What makes it different",
        [
            "Not a single-field gadget.",
            "Not only valve automation.",
            "Not only AWD digitization.",
            "It is a governance platform spanning field, cluster, and authority layers.",
        ],
        title_color=GREEN,
    )
    bullet_panel(
        s, 8.84, 1.65, 4.0, 4.95, "Evidence basis",
        [
            "IoT-assisted rice irrigation improves water-use efficiency and reduces manual dependence [5].",
            "Remote sensing expands coverage beyond instrumented plots [6].",
            "Precision irrigation becomes stronger when decisions remain auditable and supervised [7].",
        ],
        title_color=ORANGE,
    )
    add_footer(s, prs, 3, total)

    # 4 whole architecture
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_architecture.jpg")
    brand_header(s)
    add_title(s, "Whole Platform Architecture")
    add_arch_box(s, 0.4, 1.65, 2.15, 1.35, "Data Sources", ["Field sensors", "Flow / pressure", "Weather", "Crop observations"], TEAL)
    add_arch_box(s, 0.4, 3.35, 2.15, 1.35, "Remote Sensing", ["Multispectral", "Hyperspectral", "Drone / satellite", "Stress signatures"], CYAN)
    add_arch_box(s, 2.95, 2.5, 2.4, 1.45, "Edge + Gateway", ["LoRaWAN nodes", "Village gateway", "Solar + battery", "Local safety rules"], GREEN)
    add_arch_box(s, 5.65, 2.05, 2.3, 2.3, "AI Decision Engine", ["Forecast demand", "Crop-health inference", "Waterlogging risk", "Anomaly detection", "Priority scoring"], ORANGE)
    add_arch_box(s, 8.25, 1.65, 2.2, 1.35, "Governance Layer", ["Confidence score", "Policy rules", "Approval paths", "Audit trail"], TEAL_DARK)
    add_arch_box(s, 8.25, 3.35, 2.2, 1.35, "Control Layer", ["Valves", "Actuators", "Pumps / boosters", "Outlet scheduling"], TEAL_DARK)
    add_arch_box(s, 10.75, 2.5, 2.15, 1.45, "User Surfaces", ["Farmer app", "Operator console", "Authority dashboard", "Alerts"], GREEN)
    add_text(s, 2.45, 2.98, 0.38, 0.3, ["→"], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, 5.28, 2.98, 0.3, 0.3, ["→"], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, 7.98, 2.35, 0.3, 0.3, ["→"], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, 7.98, 3.75, 0.3, 0.3, ["→"], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, 10.47, 2.98, 0.25, 0.3, ["→"], 20, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, 2.95, 5.2, 7.0, 0.7, ["Closed loop: sense → infer → decide → actuate → verify → govern"], 17, WHITE, True, PP_ALIGN.CENTER)
    add_footer(s, prs, 4, total)

    # 5 multimodal intelligence
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_sensors.jpg")
    brand_header(s)
    add_title(s, "Multimodal Intelligence: What the Platform Actually Understands")
    bullet_panel(
        s, 0.55, 1.7, 3.0, 4.95, "Soil + water signals",
        [
            "Moisture, water depth, NPK, EC, pH, nutrient status",
            "Inlet/outlet flow, pressure, delivered volume",
            "Leakage, low pressure, blocked line anomalies",
        ],
    )
    bullet_panel(
        s, 3.7, 1.7, 3.0, 4.95, "Crop + imagery signals",
        [
            "Canopy vigor and vegetation indices",
            "Stress, chlorosis, disease or pest signatures",
            "Spatial variation across many fields without putting sensors everywhere [6]",
        ],
        title_color=GREEN,
    )
    bullet_panel(
        s, 6.85, 1.7, 3.0, 4.95, "Climate + forecast signals",
        [
            "Rainfall, temperature, humidity, wind, ET0",
            "Short-term forecast-aware irrigation decisions",
            "Delay irrigation when rainfall probability is high [4][5]",
        ],
        title_color=ORANGE,
    )
    bullet_panel(
        s, 10.0, 1.7, 2.85, 4.95, "Decision outputs",
        [
            "Irrigate now / wait",
            "How much water to release",
            "Which field gets priority",
            "Alert disease / stress / waterlogging",
            "Escalate to human operator if confidence is low [7]",
        ],
        title_color=TEAL_DARK,
    )
    add_footer(s, prs, 5, total)

    # 6 governance logic
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_method.jpg")
    brand_header(s)
    add_title(s, "Governance Model: Automation with Human Supervision")
    add_panel(s, 0.7, 1.8, 12.0, 4.8, fill=WHITE, alpha=228)
    steps = [
        ("1. Sense", "Live field, network, weather, and imagery inputs"),
        ("2. Analyze", "AI estimates water need, crop condition, and risk"),
        ("3. Decide", "Recommendation + confidence + rule check"),
        ("4. Actuate", "Valve / pump / outlet command if safe"),
        ("5. Verify", "Feedback from sensors confirms what happened"),
        ("6. Govern", "Farmer, operator, and authority dashboards review / override"),
    ]
    y = 2.15
    for idx, (t, d) in enumerate(steps, start=1):
        add_panel(s, 1.0, y, 2.0, 0.55, fill=LIGHT, alpha=245)
        add_text(s, 1.1, y + 0.12, 1.8, 0.25, [t], 15, TEAL_DARK, True, PP_ALIGN.CENTER)
        add_panel(s, 3.2, y, 8.9, 0.55, fill=WHITE, alpha=250)
        add_text(s, 3.35, y + 0.12, 8.55, 0.25, [d], 13, GRAY)
        y += 0.68
    add_text(s, 1.0, 6.35, 11.9, 0.28, ["The platform is strongest when AI is accountable, explainable, and embedded inside governance workflows rather than replacing them. [7]"], 12, MUTED, False, PP_ALIGN.CENTER)
    add_footer(s, prs, 6, total)

    # 7 infrastructure slide, lighter PDN
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_infra.jpg")
    brand_header(s)
    add_title(s, "Infrastructure Layer: PDN as an Enabler, Not the Entire Story")
    bullet_panel(
        s, 0.55, 1.7, 4.0, 4.95, "Infrastructure focus",
        [
            "Main source connection, filtration, sensing, actuation, and power",
            "Village sub-main junctions with volumetric measurement and control",
            "Laterals and outlets for controlled delivery closer to the field",
        ],
    )
    bullet_panel(
        s, 4.72, 1.7, 4.0, 4.95, "Why it matters",
        [
            "PDN reduces losses and supports more measurable delivery than open canals [1][2].",
            "Pressurized or assisted flow helps where terrain is uneven [2].",
            "It creates the controllable physical layer that AI can optimize.",
        ],
        title_color=GREEN,
    )
    bullet_panel(
        s, 8.89, 1.7, 3.95, 4.95, "What not to overstate",
        [
            "PDN alone is not the innovation.",
            "Sensors alone are not the innovation.",
            "The innovation is the integrated cyber-physical governance platform built on top of them.",
        ],
        title_color=ORANGE,
    )
    add_footer(s, prs, 7, total)

    # 8 operating levels
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_architecture.jpg")
    brand_header(s)
    add_title(s, "How It Works Across Operating Levels")
    level_data = [
        (GREEN, "Level 1: Field", "Reference plots and selected fields provide dense sensing, outlet control, crop observations, and ground truth for the models."),
        (CYAN, "Level 2: Cluster / Village", "Gateways aggregate telemetry, optimize allocation, compare demand vs supply, and flag deficit / excess zones."),
        (ORANGE, "Level 3: Governance", "Authorities monitor usage, efficiency, alerts, distribution fairness, infrastructure health, and intervention needs."),
    ]
    y = 1.75
    for color, title, body in level_data:
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(y), Inches(11.85), Inches(1.42))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        shape.adjustments[0] = 0.08
        add_text(s, 1.0, y + 0.15, 11.35, 0.32, [title], 18, WHITE, True)
        add_text(s, 1.0, y + 0.58, 11.35, 0.55, [body], 13, WHITE)
        y += 1.62
    add_footer(s, prs, 8, total)

    # 9 why better than conventional
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_problem.jpg")
    brand_header(s)
    add_title(s, "Why This Is Stronger Than Conventional Irrigation Alone")
    add_panel(s, 0.55, 1.7, 12.2, 4.95, fill=WHITE, alpha=230)
    add_text(s, 0.8, 1.95, 2.6, 0.3, ["Approach"], 13, TEAL, True)
    add_text(s, 3.35, 1.95, 4.2, 0.3, ["Strength"], 13, TEAL, True)
    add_text(s, 7.7, 1.95, 4.7, 0.3, ["Limitation"], 13, TEAL, True)
    rows = [
        ("Basin / continuous flooding", "Simple and familiar", "High water use, weak timing control, higher waterlogging risk [3][4]"),
        ("Manual AWD", "Saves water compared with flooding", "Still depends on frequent manual observation [4][5]"),
        ("Basic valve automation", "Reduces manual operation", "Does not understand crop health, weather, or risk"),
        ("AI + IoT + imagery + governance", "Forecast-aware, auditable, scalable, and field-sensitive", "Higher setup complexity, but much stronger operational intelligence [5][6][7]"),
    ]
    y = 2.45
    for a, b, c in rows:
        add_text(s, 0.8, y, 2.35, 0.6, [a], 12, TEAL_DARK, True)
        add_text(s, 3.35, y, 4.0, 0.6, [b], 12, GRAY)
        add_text(s, 7.7, y, 4.5, 0.6, [c], 12, MUTED)
        y += 0.95
    add_footer(s, prs, 9, total)

    # 10 expected outcomes
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_ai.jpg")
    brand_header(s)
    add_title(s, "Expected Outcomes")
    outcomes = [
        ("Water efficiency", "Lower conveyance and over-irrigation loss [1][2][3]"),
        ("Crop health visibility", "Earlier detection of stress, disease, and nutrient problems [6]"),
        ("Operational transparency", "Measure what was allocated, delivered, and consumed [7]"),
        ("Scalability", "Use dense sensors only in selected fields; extend with imagery elsewhere [6]"),
        ("Human safety", "Keep override, approval, and audit mechanisms at every level [7]"),
        ("Policy value", "Transform irrigation from infrastructure management into evidence-based governance"),
    ]
    positions = [(0.65, 1.8), (4.45, 1.8), (8.25, 1.8), (0.65, 4.2), (4.45, 4.2), (8.25, 4.2)]
    for (title, body), (x, y) in zip(outcomes, positions):
        bullet_panel(s, x, y, 3.45, 1.85, title, [body], bullet_size=12)
    add_footer(s, prs, 10, total)

    # 11 references
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_title_soft.jpg")
    brand_header(s)
    add_title(s, "Numbered References", size=26)
    add_panel(s, 0.45, 1.55, 12.45, 5.5, fill=WHITE, alpha=236)
    add_text(
        s, 0.7, 1.85, 11.95, 4.9,
        [
            "[1] Hydrospatial modelling / FAO-Aquastat aligned canal conveyance loss literature for irrigation systems in India.",
            "[2] PDN vs CDN efficiency and piped irrigation feasibility studies; CWC-aligned piped network efficiency framing.",
            "[3] Rice water use in cultivation literature: irrigation share, deep percolation, puddling and seepage losses.",
            "[4] Water-saving rice irrigation literature covering AWD, basin flooding limits, and climate-linked irrigation constraints.",
            "[5] IoT and automated AWD studies in rice systems showing added water-use efficiency over manual operation.",
            "[6] Multispectral / hyperspectral and remote-sensing literature for crop stress, health, and disease / pest indicators.",
            "[7] Precision irrigation, cyber-physical control, and human-in-the-loop governance concepts for safe deployment.",
        ],
        12, GRAY
    )
    add_footer(s, prs, 11, total)

    # 12 close
    s = prs.slides.add_slide(blank)
    add_bg(s, prs, "bg_title_soft.jpg")
    brand_header(s)
    add_panel(s, 1.8, 1.8, 9.75, 4.25, fill=WHITE, alpha=240)
    add_text(s, 2.1, 2.3, 9.15, 0.55, ["Thank You"], 34, TEAL_DARK, True, PP_ALIGN.CENTER)
    add_text(s, 2.1, 3.05, 9.15, 0.35, ["Questions & Discussion"], 19, TEAL, False, PP_ALIGN.CENTER)
    add_text(
        s, 2.1, 3.8, 9.15, 1.35,
        [
            "Arsalan Firdous",
            "SKUAST-Kashmir",
            "AI + IoT + Remote Sensing + Intelligent Irrigation Governance",
        ],
        15, GRAY, False, PP_ALIGN.CENTER
    )
    add_footer(s, prs, 12, total)

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build_presentation()
